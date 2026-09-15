from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "submission"
FIGURE_WIDTH = Inches(6.7)


def citation_numbers(text: str) -> str:
    def repl(match):
        ids = [x for x in match.group(1).split(",") if x.startswith("E")]
        numbers = {f"E{i:03d}": n for n, i in enumerate(
            [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 18], start=1
        )}
        return "[" + ", ".join(str(numbers[x]) for x in ids) + "]"

    return re.sub(r"\[((?:E\d{3},?)+)\]", repl, text)


def clean_text(text: str, language: str) -> str:
    text = re.sub(r"\[claim:C\d{3}\]\s*", "", text)
    text = re.sub(r"\[evidence:E\d{3}(?:,\s*E\d{3})*\]\s*", "", text)
    text = citation_numbers(text)
    text = text.replace("`", "")
    text = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\boldsymbol\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\text\{([^{}]+)\}", r"\1", text)
    text = text.replace(r"\,", " ").replace(r"\cdot", "·")
    text = text.replace(r"\ldots", "…")
    text = re.sub(r"\\\((.*?)\\\)", r"\1", text)
    text = re.sub(r"\[(?:to be completed|待补)[^\]]*\]", "当前记录未提供" if language == "cn" else "not available in the current records", text, flags=re.I)
    text = text.replace("（待补）", "（当前记录未提供）" if language == "cn" else "(not available in the current records)")
    text = text.replace("(to be completed)", "(not available in the current records)")
    text = text.replace("（待作者确认）", "（作者须在投稿前确认）")
    if language == "cn":
        text = re.sub(r"待补/未找到", "当前记录未提供/未找到", text)
        text = re.sub(r"(?<!补)待补(?!齐)", "当前记录未提供", text)
    text = text.replace("current results draft", "current results")
    text = text.replace("currently supportable version", "supportable results")
    text = text.replace("working-draft reference list", "numbered reference list")
    text = text.replace("working draft reference list", "numbered reference list")
    text = re.sub(r"^(\d+)\.\s+E\d{3}\.\s+", r"\1. ", text)
    math_replacements = [
        (r"\\boldsymbol\{\\theta\}", "θ"), (r"\\boldsymbol\{\\sigma\}", "σ"),
        (r"\\boldsymbol\{\\varepsilon\}", "ε"), (r"\\boldsymbol\{u\}", "u"),
        (r"\\boldsymbol\{t\}", "t"), (r"\\boldsymbol", ""), (r"\\mathbf\s*", ""),
        (r"\\mathcal\s*", ""), (r"\\mathrm\s*", ""),
        (r"\\bar\s*", "bar "), (r"\\hat\s*", "hat "),
        (r"\\theta", "θ"), (r"\\sigma", "σ"), (r"\\varepsilon", "ε"),
        (r"\\partial", "∂"), (r"\\infty", "∞"), (r"\\int", "∫"),
        (r"\\sum", "Σ"), (r"\\arg", "arg"), (r"\\min", "min"),
        (r"\\in", "∈"), (r"\\cdot", "·"), (r"\\,", " "),
        (r"\\left", ""), (r"\\right", ""),
    ]
    for pattern, replacement in math_replacements:
        text = re.sub(pattern, replacement, text)
    text = text.replace("{", "").replace("}", "")
    if language == "cn":
        if text.startswith("**图1**"):
            return "**图1** 可追溯单/双轴测量链示意图。该图为方法流程示意，不包含项目实验数值。"
        if text.startswith("**图2**"):
            return text.replace("当前不具备硬件触发验证。", "未进行硬件触发验证。")
        if text.startswith("**图3**"):
            return text.replace("所有标记均需人工复核，", "所有标记均需复核，")
        if "本章中的方括号字段必须" in text:
            return "材料、试样、设备和 DIC 元数据按原始记录填写；未找到的字段保留为当前记录未提供。"
        if text.startswith("当前版本：Data availability"):
            return "项目数据和分析输出保存在原始实验工作区；公开或受限共享方式须由作者依据实验室数据管理和保密政策确认。"
        if text.startswith("参考文献条目及 DOI/题名程序化核验结果"):
            return "参考文献按目标期刊数字格式整理，DOI 保留以便核验。"
        if text.startswith("## 结论"):
            return "## 结论"
        if text.startswith("- 保密审查："):
            return "- 保密审查：投稿前由作者按目标期刊和单位要求完成。"
        if text.startswith("- 利益冲突："):
            return "- 利益冲突：由全部作者在投稿前确认并声明。"
        if text.startswith("- 基金："):
            return "- 基金：由作者在投稿前确认并声明。"
        if text.startswith("- 作者贡献："):
            return "- 作者贡献：由全部作者按 CRediT 体系确认。"
        if text.startswith("- AI 使用："):
            return "- AI 使用：按目标期刊现行政策由作者确认并声明，作者对全部科学内容负责。"
    else:
        if text.startswith("**Figure 1.**"):
            return "**Figure 1.** Traceable uniaxial and biaxial measurement chain. This methodological workflow schematic contains no project-specific experimental values."
        if text.startswith("Current status: Data availability"):
            return "The project data and analysis outputs are retained in the originating laboratory workspace. Public or controlled access must be confirmed by the authors under the laboratory data-management and confidentiality policy."
        if text.startswith("The reference records and DOI/title validation"):
            return "The references are formatted in a numbered style, with DOI strings retained for verification."
        if text.startswith("- Confidentiality review:"):
            return "- Confidentiality review: to be completed by the authors before submission under the target journal and institutional requirements."
        if text.startswith("- Competing interests:"):
            return "- Competing interests: to be confirmed and declared by all authors before submission."
        if text.startswith("- Funding:"):
            return "- Funding: to be confirmed and declared by the authors before submission."
        if text.startswith("- Author contributions:"):
            return "- Author contributions: to be confirmed by all authors using the CRediT taxonomy."
        if text.startswith("- AI use:"):
            return "- AI use: to be confirmed under the target journal's current policy; the authors remain responsible for all scientific content."
    return text


def set_font(run, language: str, size=10.5, bold=False, italic=False, color="000000", name=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    if name is None:
        name = "Times New Roman" if language == "en" else "宋体"
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name if language == "cn" else "宋体")


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(cached)
    run._r.append(end)


def rich_runs(paragraph, text, language, size=10.5):
    pattern = re.compile(r"(\*\*.*?\*\*|\*.*?\*|\[[0-9, ]+\])")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, language, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, language, size=size, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, language, size=size, italic=True)
        else:
            run = paragraph.add_run(token)
            set_font(run, language, size=size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, language, size=size)


def omml_run(text):
    run = OxmlElement("m:r")
    text_node = OxmlElement("m:t")
    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = text
    run.append(text_node)
    return run


def omml_sub(base, sub):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(omml_run(base))
    s = OxmlElement("m:sub")
    s.append(omml_run(sub))
    node.extend([e, s])
    return node


def omml_sup(base, sup):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    e.append(omml_run(base))
    s = OxmlElement("m:sup")
    s.append(omml_run(sup))
    node.extend([e, s])
    return node


def omml_sub_sup(base, sub, sup):
    node = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    e.append(omml_run(base))
    s = OxmlElement("m:sub")
    s.append(omml_run(sub))
    sup_node = OxmlElement("m:sup")
    sup_node.append(omml_run(sup))
    node.extend([e, s, sup_node])
    return node


def omml_eq(index):
    """Build the five displayed equations used by both manuscript versions."""
    nodes = []
    if index == 1:
        nodes.extend([omml_run("g = ("), omml_sub("h", "c"), omml_run(", "),
                      omml_sub("N", "s"), omml_run(", "), omml_sub("l", "s"), omml_run(", "),
                      omml_sub("r", "f"), omml_run(", "), omml_sub("s", "s"), omml_run(", "),
                      omml_sub("w", "a"), omml_run(", …)" )])
    elif index == 2:
        nodes.extend([omml_sub("t", "m"), omml_run("(k) = a + b k")])
    elif index == 3:
        nodes.extend([omml_sub("e", "i"), omml_run(" = "),
                      omml_sub_sup("t", "m,i", "mapped"), omml_run(" − "),
                      omml_sub_sup("t", "m,i", "observed")])
    elif index == 4:
        nodes.extend([omml_sub("r", "k"), omml_run("(θ) = "), omml_sub("∫", "V"),
                      omml_run(" σ(ε(x, "), omml_sub("t", "k"), omml_run("); θ) : ε*(x) dV − "),
                      omml_sub("∫", "∂V"), omml_run(" t̄ · u*(x) dA")])
    elif index == 5:
        nodes.extend([omml_run("θ̂ = arg min"), omml_sub("θ", "θ"), omml_run(" "),
                      omml_sub("Σ", "k∈F"), omml_run(" "), omml_sub("Σ", "q∈Q"),
                      omml_run(" w"), omml_sub_sup("r", "kq", "2"),
                      omml_run("(θ)")])
    return nodes


def add_omml_equation(paragraph, index, language):
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    for node in omml_eq(index):
        math.append(node)
    math_para.append(math)
    paragraph._p.append(math_para)
    number = paragraph.add_run(f"  ({index})")
    set_font(number, language, size=9.5)


def table_widths(n):
    options = {
        4: [1.1, 2.2, 2.1, 1.6],
        5: [0.95, 1.05, 1.65, 1.45, 1.9],
        6: [1.0, 1.1, 1.55, 1.15, 1.25, 0.95],
        9: [0.8, 0.55, 0.65, 0.65, 0.75, 0.9, 0.75, 0.75, 1.2],
    }
    values = options.get(n, [7.0 / n] * n)
    scale = 7.0 / sum(values)
    return [Inches(v * scale) for v in values]


def add_table(doc, rows, language):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_widths(cols)
    for i, row in enumerate(rows):
        tr = table.rows[i]
        tr_pr = tr._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if i == 0:
            repeat_header(tr)
        for j, value in enumerate(row):
            cell = tr.cells[j]
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if i == 0:
                shade_cell(cell, "1F4E79")
            elif i % 2 == 0:
                shade_cell(cell, "F3F6F9")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or re.fullmatch(r"[-+0-9.,/ ]+", value) else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            rich_runs(paragraph, value, language, size=8.0 if cols >= 8 else 8.5)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_markdown(path, language):
    lines = path.read_text(encoding="utf-8").splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i, line in enumerate(lines[1:], start=1) if line.strip() == "---")
        lines = lines[end + 1 :]
    result = []
    i = 0
    equation_index = 0
    in_code = False
    code = []
    while i < len(lines):
        raw = lines[i]
        if raw.startswith("```"):
            if in_code:
                result.append(("code", "\n".join(code)))
                code = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(raw)
            i += 1
            continue
        if raw.strip() in (r"\[", r"\]"):
            if raw.strip() == r"\[":
                eq = []
                i += 1
                while i < len(lines) and lines[i].strip() != r"\]":
                    eq.append(lines[i].strip())
                    i += 1
                equation_index += 1
                result.append(("equation", (equation_index, " ".join(eq))))
            i += 1
            continue
        if raw.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x) for x in cells):
                    rows.append(cells)
                i += 1
            if rows:
                result.append(("table", rows))
            continue
        if raw.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", raw)
            if match:
                result.append(("image", (match.group(1), match.group(2))))
            i += 1
            continue
        if not raw.strip():
            i += 1
            continue
        result.append(("line", raw))
        i += 1
    return result


def add_line(doc, raw, language, in_references=False):
    text = clean_text(raw, language).strip()
    if not text:
        return
    if "**Evidence for" in text or text.startswith("**证据"):
        return
    if text.startswith("> "):
        text = text[2:]
    if text.startswith("**Authors:**") or text.startswith("**Funding:**") or text.startswith("**作者：**") or text.startswith("**基金：**"):
        return
    if text.startswith("## ") or text.startswith("### ") or text.startswith("# "):
        level = len(text) - len(text.lstrip("#"))
        heading = text[level:].strip()
        if language == "en" and heading.startswith("Abstract"):
            heading = "Abstract"
        if language == "cn" and heading.startswith("摘要"):
            heading = "摘要"
        heading = re.sub(r"\s*[（(](?:当前可确认版本|currently supportable version|supportable results)[）)]", "", heading)
        style = "Title" if level == 1 else ("Heading 1" if level == 2 else "Heading 2")
        p = doc.add_paragraph(style=style)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(heading)
        set_font(run, language, size=16 if level == 1 else (12.5 if level == 2 else 11), bold=True)
        return
    if text.startswith("**") and text.endswith("**") and len(text) > 4:
        text = text[2:-2]
    if in_references and re.match(r"^\d+\.\s+", text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        rich_runs(p, text, language, size=9.2)
        return
    ordered = re.match(r"^\d+\.\s+(.*)$", text)
    unordered = re.match(r"^[-*]\s+(.*)$", text)
    if ordered:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        rich_runs(p, ordered.group(0), language)
        return
    if unordered:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        rich_runs(p, unordered.group(1), language)
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    if text.startswith("**关键词：**") or text.startswith("**Keywords:**"):
        rich_runs(p, text, language, size=10)
    else:
        rich_runs(p, text, language)


def build(source_name, language, output_name):
    source = ROOT / source_name
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman" if language == "en" else "宋体"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if language == "cn" else "Times New Roman")
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Times New Roman" if language == "en" else "宋体"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if language == "cn" else "Times New Roman")
    if "PaperTitle" in styles:
        title_style = styles["PaperTitle"]
    else:
        title_style = styles.add_style("PaperTitle", WD_STYLE_TYPE.PARAGRAPH)
    title_style.base_style = normal
    title_style.font.name = "Times New Roman" if language == "en" else "宋体"
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if language == "cn" else "Times New Roman")
    title = "PA12单/双轴拉伸中的图像—载荷时序可追溯性审计" if language == "cn" else "Traceable Image–Load Registration in Uniaxial and Biaxial PA12 Tests"
    title_p = doc.add_paragraph(style="PaperTitle")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(9)
    run = title_p.add_run(title)
    set_font(run, language, size=16, bold=True)
    state = "投稿准备版；作者、基金、声明及尚未获得的实验元数据须在投稿前由作者补齐" if language == "cn" else "Submission-preparation copy; author details, declarations, and unavailable experimental metadata must be completed before submission"
    state_p = doc.add_paragraph()
    state_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    state_p.paragraph_format.space_after = Pt(12)
    state_run = state_p.add_run(state)
    set_font(state_run, language, size=9, italic=True, color="666666")
    in_references = False
    first_h1_skipped = False
    for kind, payload in parse_markdown(source, language):
        if kind == "line":
            cleaned = clean_text(payload, language)
            if not first_h1_skipped and cleaned.startswith("# "):
                first_h1_skipped = True
                continue
            if re.match(r"^(##\s+)?(参考文献|References)", cleaned):
                in_references = True
            add_line(doc, payload, language, in_references)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.space_after = Pt(5)
            for idx, line in enumerate(payload.splitlines()):
                if idx:
                    p.add_run().add_break()
                run = p.add_run(line)
                set_font(run, language, size=8.5, name="Consolas")
        elif kind == "equation":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            add_omml_equation(p, payload[0], language)
        elif kind == "table":
            add_table(doc, [[clean_text(c, language) for c in row] for row in payload], language)
        elif kind == "image":
            alt, rel = payload
            image = (source.parent / rel).resolve()
            if image.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(4)
                p.add_run().add_picture(str(image), width=FIGURE_WIDTH)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(3)
    page_number(footer)
    OUT.mkdir(parents=True, exist_ok=True)
    output = OUT / output_name
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build("manuscript_cn.md", "cn", "manuscript_cn_submission_preparation.docx")
    build("manuscript_en.md", "en", "manuscript_en_strain_submission_preparation.docx")
